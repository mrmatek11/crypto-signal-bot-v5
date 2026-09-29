"""Kanoniczny model danych importu i wspólne narzędzia parserów."""

from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from typing import Dict, Iterable, List, Optional, Sequence
from zoneinfo import ZoneInfo


@dataclass(frozen=True)
class Fill:
    """Jedno wykonanie zlecenia — źródło prawdy dla całego silnika pozycji."""

    external_id: str
    ts: datetime                      # zawsze UTC
    symbol: str
    side: str                         # "buy" | "sell"
    qty: Decimal                      # w jednostkach brokera (np. loty)
    price: Decimal
    contract_size: Decimal = Decimal(1)  # ile jednostek bazowych w 1 qty (złoto: 100 oz / lot)
    fee: Decimal = Decimal(0)         # prowizja + swap + inne koszty (ujemne = koszt)
    broker_pnl: Optional[Decimal] = None  # zrealizowany PnL raportowany przez brokera (dla zamknięć)
    stop_loss: Optional[Decimal] = None
    currency: str = "USD"


@dataclass(frozen=True)
class CashFlow:
    """Wpłata (+) lub wypłata (−) — potrzebna do stóp zwrotu ważonych czasem."""

    external_id: str
    ts: datetime                      # UTC
    amount: Decimal
    currency: str = "USD"
    note: str = ""


@dataclass
class ImportResult:
    fills: List[Fill] = field(default_factory=list)
    errors: List[str] = field(default_factory=list)
    detected: str = ""
    cash_flows: List[CashFlow] = field(default_factory=list)


class ImportErrorWithRow(ValueError):
    pass


# ─── Parsowanie wartości ────────────────────────────────────────────────────────

def norm_header(name: str) -> str:
    """„Open time ” → „open_time”, „Gross P/L” → „gross_p_l”."""
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def to_decimal(value, *, allow_empty: bool = False) -> Optional[Decimal]:
    """Liczba z formatów brokerów: '1 234,56', '1,234.56', '-0.00', '' → Decimal."""
    if value is None:
        if allow_empty:
            return None
        raise ValueError("pusta wartość liczbowa")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    s = str(value).strip().replace(" ", "").replace(" ", "")
    if s in ("", "-", "—"):
        if allow_empty:
            return None
        raise ValueError("pusta wartość liczbowa")
    if "," in s and "." in s:
        # separator dziesiętny to ten, który występuje później
        if s.rfind(",") > s.rfind("."):
            s = s.replace(".", "").replace(",", ".")
        else:
            s = s.replace(",", "")
    elif "," in s:
        s = s.replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise ValueError(f"niepoprawna liczba: {value!r}") from exc


DATE_FORMATS = (
    "%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y.%m.%d %H:%M:%S", "%Y.%m.%d %H:%M",
    "%d.%m.%Y %H:%M:%S", "%d.%m.%Y %H:%M", "%d/%m/%Y %H:%M:%S", "%Y-%m-%dT%H:%M:%S",
)


def to_utc(value, tz: str = "UTC", fmt: Optional[str] = None) -> datetime:
    """Czas z pliku (w strefie `tz`) → UTC. Obsługuje datetime z XLSX i tekst."""
    if isinstance(value, datetime):
        dt = value
    else:
        s = str(value).strip()
        dt = None
        for f in ((fmt,) if fmt else DATE_FORMATS):
            try:
                dt = datetime.strptime(s, f)
                break
            except ValueError:
                continue
        if dt is None:
            try:
                dt = datetime.fromisoformat(s)
            except ValueError as exc:
                raise ValueError(f"nieznany format daty: {value!r}") from exc
    if dt.tzinfo is None:
        dt = dt.replace(tzinfo=ZoneInfo(tz))
    return dt.astimezone(timezone.utc)


# ─── Czytanie plików ───────────────────────────────────────────────────────────

def read_rows(data: bytes, filename: str) -> List[Dict[str, object]]:
    """CSV (auto-separator) albo XLSX → lista słowników z znormalizowanymi nagłówkami.

    Dla XLSX szuka wiersza nagłówka (pierwszy wiersz z ≥ 4 niepustymi komórkami tekstowymi),
    bo eksporty brokerów często mają nad tabelą nagłówek raportu.
    """
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        best: List[Dict[str, object]] = []
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            header_idx = next((i for i, r in enumerate(rows)
                               if sum(isinstance(c, str) and c.strip() != "" for c in r) >= 4), None)
            if header_idx is None:
                continue
            header = [norm_header(c) if c is not None else "" for c in rows[header_idx]]
            parsed = [
                {h: v for h, v in zip(header, r) if h}
                for r in rows[header_idx + 1:]
                if r and any(v not in (None, "") for v in r)
            ]
            if len(parsed) > len(best):
                best = parsed
        return best

    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        return []
    header = [norm_header(c) for c in rows[0]]
    return [{h: v for h, v in zip(header, r) if h} for r in rows[1:]]


def pick(row: Dict[str, object], *names: str):
    """Pierwsza istniejąca kolumna z listy aliasów."""
    for n in names:
        if n in row and row[n] not in (None, ""):
            return row[n]
    return None


def require(row: Dict[str, object], *names: str):
    v = pick(row, *names)
    if v is None:
        raise ValueError(f"brak kolumny/wartości: {' | '.join(names)}")
    return v


def has_columns(rows: Sequence[Dict[str, object]], *names: str) -> bool:
    return bool(rows) and all(any(n in r for r in rows[:5]) for n in names)


def iter_with_errors(rows: Iterable[Dict[str, object]], fn, result: ImportResult, first_line: int = 2):
    """Uruchom parser wiersza; błędy zbieraj z numerem wiersza zamiast przerywać cały import."""
    for i, row in enumerate(rows):
        try:
            result.fills.extend(fn(row))
        except (ValueError, KeyError) as exc:
            result.errors.append(f"wiersz {i + first_line}: {exc}")
