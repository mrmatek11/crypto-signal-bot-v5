"""cTrader — eksport historii zamkniętych pozycji (zakładka History → Export, CSV/XLSX).

Każdy wiersz to zamknięta pozycja → dwa fill-e. Kolumny (nazwy dopasowujemy po aliasach):
ID / Position ID, Symbol, Opening Direction, Opening Time, Closing Time, Entry Price, Closing Price,
Closing Quantity (np. „0.50 Lots” albo „50 Oz”), Commission, Swap, Gross USD / Net USD.

Czas jest w strefie wybranej w cTrader; gdy nagłówek ją zawiera („Opening Time (UTC+2)”), używamy jej,
w przeciwnym razie `tz` (domyślnie UTC). Ilość w uncjach/jednostkach zapisujemy z wielkością kontraktu 1.
⚠️ Nazwy kolumn zweryfikować na prawdziwym eksporcie od brokera (różnią się między wersjami cTrader).
"""

from __future__ import annotations

import re
from datetime import timedelta
from decimal import Decimal
from typing import Dict, List, Optional, Tuple

from .base import Fill, ImportResult, iter_with_errors, pick, read_rows, require, to_decimal, to_utc
from .instruments import default_contract_size, infer_contract_size, normalize_symbol

NAME = "ctrader"
_UTC_SUFFIX = re.compile(r"_utc_?([0-9]{1,2})(?:_([0-9]{2}))?$")
_QTY = re.compile(r"^\s*([-+]?[0-9][0-9\s.,]*)\s*([A-Za-z()]*)\s*$")


def _col(row_keys, *prefixes: str) -> Optional[str]:
    for p in prefixes:
        for k in row_keys:
            if k == p or k.startswith(p + "_"):
                return k
    return None


def detect(rows) -> bool:
    if not rows:
        return False
    keys = set().union(*(r.keys() for r in rows[:5]))
    return bool(_col(keys, "opening_direction")) or (
        bool(_col(keys, "entry_price")) and bool(_col(keys, "closing_price")) and bool(_col(keys, "closing_time")))


def _offset(key: str, sign_hint: str) -> Optional[timedelta]:
    """„opening_time_utc_2” → +2 h. Znak minus znika przy normalizacji nagłówka, więc bierzemy go z surowej nazwy."""
    m = _UTC_SUFFIX.search(key)
    if not m:
        return None
    delta = timedelta(hours=int(m.group(1)), minutes=int(m.group(2) or 0))
    return -delta if sign_hint == "-" else delta


def _time(value, key: str, tz: str, signs: Dict[str, str]):
    off = _offset(key, signs.get(key, "+"))
    if off is None:
        return to_utc(value, tz)
    return to_utc(value, "UTC") - off


def _quantity(raw, symbol: str) -> Tuple[Decimal, Optional[Decimal]]:
    """(ilość, wielkość kontraktu jeśli wynika z jednostki)."""
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw)), None
    m = _QTY.match(str(raw))
    if not m:
        raise ValueError(f"nieczytelna ilość: {raw!r}")
    qty = to_decimal(m.group(1))
    unit = m.group(2).lower()
    if unit.startswith("lot") or unit == "":
        return qty, None
    if unit in ("oz", "ounce", "ounces", "units", "unit", "u"):
        return qty, Decimal(1)
    raise ValueError(f"nieznana jednostka ilości: {m.group(2)}")


def parse(data: bytes, filename: str, tz: str = "UTC", raw_headers: Optional[Dict[str, str]] = None) -> ImportResult:
    rows = read_rows(data, filename)
    result = ImportResult(detected=NAME)
    if not detect(rows):
        result.errors.append("To nie wygląda na historię pozycji z cTrader (brak Opening Direction / Entry Price).")
        return result
    signs = raw_headers if raw_headers is not None else _header_signs(data, filename)
    keys = set().union(*(r.keys() for r in rows[:5]))
    c_open = _col(keys, "opening_time", "open_time")
    c_close = _col(keys, "closing_time", "close_time")
    c_gross = _col(keys, "gross")
    c_net = _col(keys, "net")
    if not c_open or not c_close:
        result.errors.append("Brak kolumn czasu otwarcia / zamknięcia.")
        return result

    def row_to_fills(row: Dict[str, object]) -> List[Fill]:
        pos_id = str(require(row, "id", "position_id", "position")).split(".")[0]
        symbol = normalize_symbol(require(row, "symbol"))
        side_raw = str(require(row, "opening_direction", "direction", "side", "type")).strip().lower()
        if side_raw in ("buy", "long"):
            direction = 1
        elif side_raw in ("sell", "short"):
            direction = -1
        else:
            raise ValueError(f"nieznany kierunek: {side_raw}")
        qty, unit_size = _quantity(require(row, "closing_quantity", "quantity", "volume", "lots"), symbol)
        open_px = to_decimal(require(row, "entry_price", "opening_price"))
        close_px = to_decimal(require(row, "closing_price"))
        costs = sum((to_decimal(pick(row, c), allow_empty=True) or Decimal(0)) for c in ("commission", "commissions", "swap"))
        gross = to_decimal(row.get(c_gross), allow_empty=True) if c_gross else None
        if gross is None and c_net:
            net = to_decimal(row.get(c_net), allow_empty=True)
            gross = None if net is None else net - costs
        size = unit_size or infer_contract_size(open_px, close_px, qty, direction, gross) or default_contract_size(symbol)
        open_side, close_side = ("buy", "sell") if direction == 1 else ("sell", "buy")
        return [
            Fill(external_id=f"{pos_id}:open", ts=_time(require(row, c_open), c_open, tz, signs), symbol=symbol,
                 side=open_side, qty=qty, price=open_px, contract_size=size),
            Fill(external_id=f"{pos_id}:close", ts=_time(require(row, c_close), c_close, tz, signs), symbol=symbol,
                 side=close_side, qty=qty, price=close_px, contract_size=size, fee=costs, broker_pnl=gross),
        ]

    iter_with_errors(rows, row_to_fills, result)
    return result


def _header_signs(data: bytes, filename: str) -> Dict[str, str]:
    """Znak przesunięcia UTC z surowych nagłówków (normalizacja zamienia „UTC-5” i „UTC+5” na to samo)."""
    from .base import norm_header

    text = ""
    if filename.lower().endswith((".xlsx", ".xlsm")):
        from io import BytesIO

        from openpyxl import load_workbook

        wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
        for ws in wb.worksheets:
            for r in ws.iter_rows(values_only=True, max_row=30):
                text += " | ".join(str(c) for c in r if c is not None) + "\n"
    else:
        text = data[:8192].decode("utf-8-sig", errors="replace")
    out: Dict[str, str] = {}
    for m in re.finditer(r"([A-Za-z ]*Time\s*\(?\s*UTC\s*([+-])\s*[0-9:]+\s*\)?)", text):
        out[norm_header(m.group(1))] = m.group(2)
    return out
